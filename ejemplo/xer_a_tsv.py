#!/usr/bin/env python3
"""
xer_a_tsv.py - extrae una tabla de un export de Primavera P6 en xlsx a TSV

El programa cronograma.c lee TSV, que es el formato nativo de un XER. Este
script hace la conversion desde el xlsx que exporta P6.

    python3 xer_a_tsv.py Cronograma_Contratista_DB.xlsx           -> task.tsv
    python3 xer_a_tsv.py Cronograma_Contratista_DB.xlsx TASKPRED  -> taskpred.tsv
    python3 xer_a_tsv.py archivo.xlsx TASK salida.tsv

Requiere openpyxl:  pip install openpyxl
"""

import sys
import warnings

warnings.filterwarnings("ignore")

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("falta openpyxl. Instalalo con: pip install openpyxl")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__.strip())

    ruta = sys.argv[1]
    hoja = sys.argv[2] if len(sys.argv) > 2 else "TASK"
    salida = sys.argv[3] if len(sys.argv) > 3 else hoja.lower() + ".tsv"

    wb = load_workbook(ruta, read_only=True)

    if hoja not in wb.sheetnames:
        sys.exit(f"la hoja '{hoja}' no existe. Hay: {', '.join(wb.sheetnames)}")

    ws = wb[hoja]
    it = ws.iter_rows(values_only=True)

    # Primera fila: nombres tecnicos (task_code, wbs_id, ...). Es la que usa
    # cronograma.c para ubicar las columnas por nombre.
    # Segunda fila: los titulos que muestra P6 ("Activity ID", ...). Se descarta.
    cabecera = next(it)
    next(it, None)

    cols = [i for i, c in enumerate(cabecera) if c]

    n = 0
    with open(salida, "w", encoding="utf-8") as f:
        f.write("\t".join(str(cabecera[i]) for i in cols) + "\n")
        for fila in it:
            if fila[cols[0]] is None:
                continue
            f.write("\t".join(
                "" if fila[i] is None else str(fila[i]).replace("\t", " ").replace("\n", " ")
                for i in cols
            ) + "\n")
            n += 1

    print(f"{salida}: {n} filas, {len(cols)} columnas")


if __name__ == "__main__":
    main()
